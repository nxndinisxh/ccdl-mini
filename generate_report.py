import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_report(output_path):
    doc = docx.Document()

    # Page Margins: 1 inch everywhere
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style Configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 51, 51) # Charcoal
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)    # Slate 900
    COLOR_SECONDARY = RGBColor(37, 99, 235) # Blue 600
    COLOR_ACCENT = RGBColor(79, 70, 229)    # Indigo 600
    COLOR_MUTED = RGBColor(100, 116, 139)   # Slate 500

    # ==========================
    # TITLE & METADATA BLOCK
    # ==========================
    p_course = doc.add_paragraph()
    r_course = p_course.add_run("MANIPAL INSTITUTE OF TECHNOLOGY")
    r_course.bold = True
    r_course.font.size = Pt(12)
    r_course.font.color.rgb = COLOR_SECONDARY
    p_course.paragraph_format.space_after = Pt(2)

    p_dept = doc.add_paragraph()
    r_dept = p_dept.add_run("DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\nCourse: Cloud Computing & DevOps (Mini-Project) | 5th Semester | Batch D2")
    r_dept.font.size = Pt(10)
    r_dept.font.color.rgb = COLOR_MUTED
    p_dept.paragraph_format.space_after = Pt(16)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run("Cloud-Native Market Regime Detection &\nAdaptive Portfolio Risk Engine")
    r_title.bold = True
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = COLOR_PRIMARY
    p_title.paragraph_format.space_after = Pt(8)

    p_subtitle = doc.add_paragraph()
    r_sub = p_subtitle.add_run("Technical Design & Final Project Report (Lab 12 Deliverable)")
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = COLOR_SECONDARY
    p_subtitle.paragraph_format.space_after = Pt(18)

    # Metadata Table
    table_meta = doc.add_table(rows=4, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False

    meta_data = [
        ("Team Members:", "Nandini Sah (Reg: 240905552, Roll: 54, CSE-D)\nPrisha Gupta (Reg: 240905612, Roll: 61, CSE-D)"),
        ("Faculty Supervisors:", "Dr. Anup Bhat & Dr. Swati B P"),
        ("Source Code Repository:", "https://github.com/nxndinisxh/ccdl-mini.git (Branch: main)"),
        ("Deployment Target & Account:", "Amazon Web Services (AWS us-east-1) | Account ID: 777910569361")
    ]

    col_widths = [Inches(2.2), Inches(4.3)]
    for i, (label, val) in enumerate(meta_data):
        row = table_meta.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = col_widths[0], col_widths[1]
        
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c1, top=80, bottom=80, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Helper function for headings
    def add_custom_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        if level == 1:
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            r.font.size = Pt(15)
            r.font.color.rgb = COLOR_PRIMARY
        elif level == 2:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r.font.size = Pt(12)
            r.font.color.rgb = COLOR_SECONDARY
        elif level == 3:
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r.font.size = Pt(11)
            r.font.color.rgb = COLOR_ACCENT
        return p

    # ==========================
    # 1. EXECUTIVE SUMMARY & OBJECTIVES
    # ==========================
    add_custom_heading("1. Executive Summary & Project Objectives")
    
    p = doc.add_paragraph()
    p.add_run("What does this project do in simple terms? ").bold = True
    p.add_run(
        "Financial markets are unpredictable: one month they are steadily rising (bull market), and the next month they crash "
        "or fluctuate wildly. Most automated trading or portfolio tools assume market conditions stay the same, which leads to "
        "massive losses during market downturns. Our project, "
    )
    p.add_run("Cloud-Native Market Regime Detection and Adaptive Portfolio Risk Engine").bold = True
    p.add_run(
        ", is an automated cloud system that continuously monitors stock market prices, identifies the underlying market state "
        "(e.g., Bullish, Neutral, or Crash/Bearish) using unsupervised Machine Learning (K-Means), and automatically adjusts "
        "investment portfolio weights (stocks, bonds, cash) to protect capital before severe losses happen."
    )

    add_custom_heading("Core Objectives:", level=2)
    objectives = [
        ("Automated Market Regime Classification: ", "Extract quantitative momentum and volatility indicators (RSI, Rolling Volatility, SMA ratios) and use an unsupervised K-Means ML algorithm to cluster market conditions into 4 actionable states."),
        ("Dynamic Risk Management & Allocation: ", "Calculate real-time financial risk metrics—Value-at-Risk (VaR at 95% and 99%), Conditional VaR (Expected Shortfall), and Maximum Drawdown—and dynamically shift asset weights between Equities, Fixed Income, and Cash."),
        ("Event-Driven Cloud-Native Architecture: ", "Decouple data ingestion, ML clustering, persistent storage, and alerting across 8+ Amazon Web Services (AWS) managed services using asynchronous messaging."),
        ("100% Infrastructure as Code (IaC): ", "Automate the entire cloud infrastructure provisioning via Terraform with zero manual console clicking, ensuring complete reproducibility and security."),
        ("Strict Free-Tier Cost Governance: ", "Guarantee a $0.00 operational cost footprint under the AWS Free Tier, enforced by an automated AWS Budget alarm triggering at $0.01."),
        ("Production CI/CD Automation: ", "Implement a continuous integration and automated testing pipeline using GitHub Actions to validate unit tests, linting, Terraform syntax, and Docker builds on every commit.")
    ]
    for prefix, desc in objectives:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_pre = bp.add_run(prefix)
        r_pre.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY
        bp.add_run(desc)

    # ==========================
    # 2. PROBLEM STATEMENT
    # ==========================
    add_custom_heading("2. Problem Statement")
    p = doc.add_paragraph(
        "Modern institutional and retail investors face three critical challenges that this project resolves:"
    )
    challenges = [
        ("Static Model Vulnerability: ", "Traditional financial models assume market returns follow a static bell curve (Gaussian distribution). During sudden liquidity crunches (e.g., 2008 or March 2020), this causes catastrophic portfolio drawdowns because models do not recognize that market volatility has fundamentally changed."),
        ("Manual & Monolithic Cloud Infrastructures: ", "Many financial engineering pipelines are executed as manual Python scripts or deployed on single monolithic virtual machines. When traffic surges or data breaks, the entire system crashes with zero observability or auto-healing."),
        ("Uncontrolled Cloud Spend: ", "Running always-on EC2 servers or high-end relational databases for financial analysis generates expensive monthly bills. A truly modern system must be serverless—scaling up when new market ticks arrive and immediately scaling down to zero idle cost.")
    ]
    for prefix, desc in challenges:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_pre = bp.add_run(prefix)
        r_pre.bold = True
        bp.add_run(desc)

    # ==========================
    # 3. CLOUD COMPUTING & SOLUTION ARCHITECTURE
    # ==========================
    add_custom_heading("3. Cloud Computing & Solution Architecture")
    p = doc.add_paragraph(
        "The architecture is designed to be fully event-driven, decoupled, and serverless on Amazon Web Services (AWS us-east-1). "
        "No continuous compute instances (EC2) are kept running, keeping the entire deployment 100% within the AWS Free Tier."
    )

    # AWS Services Table
    table_aws = doc.add_table(rows=1, cols=3)
    table_aws.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_aws.autofit = False

    headers = ["AWS Managed Service", "Role in Architecture", "Free Tier Compliance"]
    hdr_cells = table_aws.rows[0].cells
    widths = [Inches(1.8), Inches(3.2), Inches(1.5)]
    for idx, name in enumerate(headers):
        hdr_cells[idx].width = widths[idx]
        set_cell_background(hdr_cells[idx], "1E3A8A")
        set_cell_margins(hdr_cells[idx], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[idx].paragraphs[0]
        r = p.add_run(name)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    services_data = [
        ("Amazon S3", "Raw tick storage & audit bucket (`market-regime-risk-engine-data-...`). Stores ingested market CSV/JSON files with private encryption.", "5 GB storage free"),
        ("Amazon SQS + DLQ", "Asynchronous message queue (`ingestion-queue`) that buffers raw data events. Features a Dead Letter Queue with 3-retry redrive policy.", "1 Million requests/mo free"),
        ("AWS Lambda", "Serverless worker (`market-regime-risk-engine-worker`). Runs ML clustering, risk math, and writes state on SQS trigger.", "1 Million invocations/mo free"),
        ("Amazon DynamoDB", "NoSQL database table (`MarketRegimes`) storing calculated regime states, timestamps, VaR metrics, and asset weights.", "Provisioned 5 RCU / 5 WCU (Always Free)"),
        ("Amazon SNS", "Push notification alert topic (`regime-alerts`). Dispatches immediate email notifications when High-Volatility Bearish regime or Risk > 75 occurs.", "100,000 requests/mo free"),
        ("Amazon API Gateway (v2)", "HTTP API Gateway exposing REST endpoints (`ANY /{proxy+}`) with low-latency CORS proxy to backend services.", "1 Million calls/mo free"),
        ("Amazon CloudWatch", "Centralized logging and metrics stream (`/aws/market-regime/...`) monitoring execution times and errors.", "5 GB log ingestion free"),
        ("AWS Budgets", "Financial safety circuit breaker set to trigger alert if projected monthly cost exceeds $0.01.", "Free tier enabled ($0 cost)")
    ]

    for item in services_data:
        row = table_aws.add_row()
        for idx, text in enumerate(item):
            c = row.cells[idx]
            c.width = widths[idx]
            set_cell_background(c, "F8FAFC" if row._index % 2 == 0 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if idx == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Workflow description
    add_custom_heading("End-to-End Workflow & Event Pipeline:", level=2)
    steps = [
        ("1. Data Ingestion: ", "The backend ingestion worker queries real-time and historical financial data (Yahoo Finance), serializes the dataset, and uploads the raw payload to Amazon S3 for historical audit logging."),
        ("2. Asynchronous Queuing: ", "An event is published to Amazon SQS (`ingestion-queue`). This decouples ingestion from compute, guaranteeing that network delays or market spikes do not overwhelm the analytics backend."),
        ("3. Serverless ML Execution: ", "AWS Lambda automatically triggers upon reading the SQS queue. It invokes the Feature Extraction and K-Means Regime Engine."),
        ("4. State Persistence: ", "The regime prediction, volatility metrics, VaR calculations, and rebalanced asset allocation weights are written directly into Amazon DynamoDB."),
        ("5. Proactive Anomaly Alerting: ", "If the engine detects a 'High-Volatility Bearish' crash regime or an aggregate risk score > 75, it publishes a high-priority SNS alert to investor emails instantly."),
        ("6. Real-Time Institutional Dashboard: ", "The React/Vite web application queries the API Gateway and displays regime charts, portfolio rebalancing visualizers, and stress-testing simulation sandboxes.")
    ]
    for prefix, desc in steps:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_pre = bp.add_run(prefix)
        r_pre.bold = True
        bp.add_run(desc)

    # ==========================
    # 4. MACHINE LEARNING & RISK ENGINE DESIGN
    # ==========================
    add_custom_heading("4. Machine Learning & Quantitative Risk Engine")
    
    add_custom_heading("4.1 Feature Engineering Pipeline", level=2)
    p = doc.add_paragraph(
        "Rather than feeding raw stock prices directly into the model, the feature pipeline extracts five normalized technical signals:"
    )
    features = [
        ("Logarithmic Returns: ", "Measures continuous compounding price differences: r_t = ln(P_t / P_{t-1})."),
        ("21-Day Rolling Volatility: ", "Annualized standard deviation of returns measuring recent price turbulence."),
        ("Trend Ratio (SMA-50 Deviation): ", "Percentage deviation of current price from the 50-day Simple Moving Average: (P_t - SMA_50) / SMA_50."),
        ("Moving Average Spread: ", "Difference between 20-day Exponential Moving Average and 50-day Simple Moving Average."),
        ("Relative Strength Index (RSI-14): ", "Bounded momentum indicator measuring relative magnitude of recent gains versus losses (0 to 100).")
    ]
    for prefix, desc in features:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        r_pre = bp.add_run(prefix)
        r_pre.bold = True
        bp.add_run(desc)

    add_custom_heading("4.2 Unsupervised K-Means & Semantic Centroid Mapping", level=2)
    p = doc.add_paragraph(
        "Market states naturally form clusters in multi-dimensional feature space. We train an unsupervised K-Means model (K=4). "
        "Because unsupervised clustering assigns arbitrary cluster indices (0, 1, 2, 3), we implement an automated Semantic Centroid Mapper "
        "that inspects cluster return, trend, and volatility centroids to assign definitive economic meanings:"
    )

    # Allocation Table
    table_regimes = doc.add_table(rows=1, cols=5)
    table_regimes.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_regimes.autofit = False

    r_headers = ["Market Regime", "Economic Meaning", "Equities", "Bonds", "Cash / Hedges"]
    r_hdr_cells = table_regimes.rows[0].cells
    r_widths = [Inches(1.8), Inches(2.2), Inches(0.8), Inches(0.8), Inches(0.9)]
    for idx, name in enumerate(r_headers):
        r_hdr_cells[idx].width = r_widths[idx]
        set_cell_background(r_hdr_cells[idx], "0F172A")
        set_cell_margins(r_hdr_cells[idx], top=80, bottom=80, left=60, right=60)
        p = r_hdr_cells[idx].paragraphs[0]
        r = p.add_run(name)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    regimes_data = [
        ("Low-Vol Bullish", "Steady upward trend, low volatility", "80%", "15%", "5%"),
        ("High-Vol Bullish", "Strong rally with elevated swings", "55%", "30%", "15%"),
        ("Sideways / Neutral", "Range-bound, consolidation phase", "40%", "40%", "20%"),
        ("High-Vol Bearish", "Liquidation spiral, crash danger", "15%", "35%", "50%")
    ]

    for item in regimes_data:
        row = table_regimes.add_row()
        for idx, text in enumerate(item):
            c = row.cells[idx]
            c.width = r_widths[idx]
            set_cell_background(c, "F8FAFC" if row._index % 2 == 0 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=60, right=60)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if idx == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Risk Metrics
    add_custom_heading("4.3 Quantitative Risk Calculations", level=2)
    p = doc.add_paragraph(
        "In addition to regime classification, the engine calculates institutional risk metrics:\n"
        "• Value-at-Risk (VaR 95% & 99%): The maximum expected monetary loss over a 1-day horizon at 95% and 99% confidence.\n"
        "• Conditional VaR (Expected Shortfall): The average expected loss when an extreme tail-risk event exceeds the VaR threshold.\n"
        "• Maximum Drawdown (MDD): The largest peak-to-trough decline observed across the lookback horizon.\n"
        "• Composite Risk Score (0-100): An aggregated weighted score combining regime severity, rolling volatility, VaR, and drawdown."
    )

    # ==========================
    # 5. DEVOPS, IAC & AUTOMATION
    # ==========================
    add_custom_heading("5. DevOps Stack & Infrastructure as Code")
    
    add_custom_heading("5.1 Terraform Infrastructure as Code (IaC)", level=2)
    p = doc.add_paragraph(
        "All 20 AWS cloud resources are defined declaratively in Terraform (`terraform/` directory). "
        "No cloud resources were created manually in the AWS web console. The Terraform codebase includes:\n"
        "• `main.tf`: Defines S3 bucket with server-side encryption, SQS FIFO/Standard queue, Dead Letter Queue (DLQ), DynamoDB table with 5 RCU/WCU provisioned throughput, SNS topic, CloudWatch log group, and AWS Budget.\n"
        "• `lambda.tf`: Configures IAM execution roles with least-privilege policies, Lambda function package, and SQS event source mapping.\n"
        "• `api_gateway.tf`: Creates HTTP API v2 routes and proxy integrations to Lambda.\n"
        "• `variables.tf` & `outputs.tf`: Parameterizes environment variables and outputs endpoint URLs upon apply."
    )

    add_custom_heading("5.2 CI/CD Automation (GitHub Actions)", level=2)
    p = doc.add_paragraph(
        "The project repository incorporates an automated GitHub Actions pipeline (`.github/workflows/ci-cd.yml`) that triggers on every pull request and push to the `main` branch. It executes three parallel validation stages:\n"
        "1. Automated Unit & Integration Testing: Runs `pytest` across 7 test suites validating feature extraction, K-Means clustering, risk math, and API contract conformity.\n"
        "2. Infrastructure Validation: Executes `terraform fmt -check` and `terraform validate` to ensure zero syntax or schema regressions.\n"
        "3. Docker Build Verification: Validates that both backend and frontend Docker container images compile cleanly without build errors."
    )

    add_custom_heading("5.3 Containerization (Docker & Docker Compose)", level=2)
    p = doc.add_paragraph(
        "Multi-stage Dockerfiles ensure identical behavior across development and production environments:\n"
        "• Backend Dockerfile: Python 3.11-slim container running FastAPI through Uvicorn.\n"
        "• Frontend Dockerfile: Multi-stage Node.js builder serving compiled static assets through high-performance Alpine Nginx.\n"
        "• docker-compose.yml: Enables one-command local orchestration (`docker compose up`) for testing."
    )

    # ==========================
    # 6. RESULTS & VERIFICATION
    # ==========================
    add_custom_heading("6. Live Results & Cloud Verification")
    p = doc.add_paragraph(
        "The infrastructure was deployed and verified live in the `us-east-1` AWS region (Account ID: `777910569361`). "
        "The following real cloud resources were successfully provisioned and verified:"
    )

    # Outputs table
    table_res = doc.add_table(rows=1, cols=2)
    table_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_res.autofit = False

    res_headers = ["Cloud Resource Output Name", "Live AWS Provisioned Value / ARN"]
    res_hdr_cells = table_res.rows[0].cells
    res_widths = [Inches(2.5), Inches(4.0)]
    for idx, name in enumerate(res_headers):
        res_hdr_cells[idx].width = res_widths[idx]
        set_cell_background(res_hdr_cells[idx], "1E293B")
        set_cell_margins(res_hdr_cells[idx], top=80, bottom=80, left=80, right=80)
        p = res_hdr_cells[idx].paragraphs[0]
        r = p.add_run(name)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    outputs_data = [
        ("api_gateway_endpoint", "https://ex5ctqi2a6.execute-api.us-east-1.amazonaws.com"),
        ("s3_bucket_name", "market-regime-risk-engine-data-20261025160405985900000001"),
        ("dynamodb_table_name", "MarketRegimes (Provisioned 5 RCU / 5 WCU)"),
        ("sqs_queue_url", "https://sqs.us-east-1.amazonaws.com/777910569361/market-regime-risk-engine-ingestion-queue"),
        ("sns_topic_arn", "arn:aws:sns:us-east-1:777910569361:market-regime-risk-engine-regime-alerts"),
        ("lambda_worker", "market-regime-risk-engine-worker (Runtime: Python 3.11, Memory: 512 MB)"),
        ("budget_alarm_threshold", "$0.01 USD monthly safeguard (Actual spend: $0.00)")
    ]

    for item in outputs_data:
        row = table_res.add_row()
        for idx, text in enumerate(item):
            c = row.cells[idx]
            c.width = res_widths[idx]
            set_cell_background(c, "F8FAFC" if row._index % 2 == 0 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.size = Pt(9)
            if idx == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================
    # 7. LESSONS LEARNED
    # ==========================
    add_custom_heading("7. Lessons Learned")
    lessons = [
        ("Cloud Cost Governance: ", "Serverless architectures (Lambda, API Gateway, DynamoDB) offer superior cost efficiency, but configuring strict provisioned capacity and SQS long-polling is critical to avoid accidental polling surcharges."),
        ("SQS Visibility Timeout vs. Lambda Execution: ", "During initial Terraform deployment, an event source mapping failure revealed that SQS visibility timeout must strictly exceed Lambda execution timeout. Setting SQS visibility to 180s and Lambda timeout to 30s resolved all queue redrive anomalies."),
        ("Deterministic Semantic Mapping in Unsupervised ML: ", "Unsupervised clustering algorithms (K-Means) output arbitrary cluster indices (0-3). Creating an automated semantic centroid sorting algorithm was vital to mapping mathematical clusters to actionable real-world financial decisions."),
        ("Resilience via Graceful Degradation: ", "Integrating automated fallback generators in the data ingestion pipeline ensured that external rate limits from financial data providers never break downstream ML processing or live presentations.")
    ]
    for prefix, desc in lessons:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_pre = bp.add_run(prefix)
        r_pre.bold = True
        bp.add_run(desc)

    # ==========================
    # 8. INDIVIDUAL CONTRIBUTIONS
    # ==========================
    add_custom_heading("8. Individual Contributions")
    p = doc.add_paragraph(
        "In compliance with the Cloud Computing & DevOps mini-project guidelines, team responsibilities were partitioned as follows:"
    )

    add_custom_heading("Nandini Sah (Registration Number: 240905552, Roll No: 54, CSE-D)", level=2)
    nandini_tasks = [
        "Cloud Architecture & IaC: Designed and authored the complete Terraform infrastructure (`main.tf`, `lambda.tf`, `api_gateway.tf`, `variables.tf`, `outputs.tf`).",
        "AWS Cloud Provisioning: Deployed Amazon S3, SQS queues with Dead Letter Queues, SNS alerting, and DynamoDB tables under strict Free Tier limits.",
        "Cost Safeguards: Configured and deployed AWS Budget alarms at the $0.01 threshold for zero-cost operation.",
        "CI/CD Pipeline: Architected the GitHub Actions automated workflow (`ci-cd.yml`) executing tests, Terraform validation, and Docker builds.",
        "Backend API & Orchestration: Developed the FastAPI server application (`main.py`) and Docker container configuration."
    ]
    for task in nandini_tasks:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(task)

    add_custom_heading("Prisha Gupta (Registration Number: 240905612, Roll No: 61, CSE-D)", level=2)
    prisha_tasks = [
        "Feature Engineering Pipeline: Formulated and implemented mathematical indicators (`features.py`: log returns, rolling volatility, RSI, SMA ratios).",
        "Machine Learning Regime Engine: Designed the unsupervised K-Means clustering model and semantic centroid mapping heuristic (`regime_engine.py`).",
        "Portfolio Risk Engine: Implemented quantitative financial risk metrics (`risk_engine.py`: VaR 95/99%, Expected Shortfall, MDD, Sharpe, dynamic asset allocations).",
        "Frontend Quantitative UI: Developed the institutional React 18 dashboard (`src/App.jsx`) with interactive price charts, regime bands, and shock injection sandbox.",
        "Automated Test Suite: Co-developed the comprehensive Pytest unit and integration test suite (`test_engine.py`)."
    ]
    for task in prisha_tasks:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(task)

    # Save to path
    doc.save(output_path)
    print(f"Report successfully saved to: {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\Personal Data\Desktop\projects\ccdl-mini-project"
    docs_dir = os.path.join(out_dir, "docs")
    
    file1 = os.path.join(out_dir, "Cloud_Computing_DevOps_MiniProject_Report.docx")
    file2 = os.path.join(docs_dir, "Cloud_Computing_DevOps_MiniProject_Report.docx")
    
    create_report(file1)
    create_report(file2)
